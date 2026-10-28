"""Low-level OpenXML and styling primitives for AS-IS Word DOCX annex compilation."""

from __future__ import annotations

import html
import json
import logging
import re
from pathlib import Path
from typing import Any

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from assessment_engine.infrastructure.docx_render_utils import (
    set_cell_text,
    shade_cell,
)

logger = logging.getLogger(__name__)

COLOR_BLUE = "0072BC"
COLOR_HEADER_BG = "D9EAF7"
COLOR_ROW_ALT = "F2F2F2"


def format_currency_custom(value: float, thousands_sep: str, decimal_sep: str) -> str:
    """Formats a float into a string with custom thousands and decimal separators."""
    parts = f"{value:.2f}".split(".")
    integer_part = parts[0]
    decimal_part = parts[1]
    reversed_integer = integer_part[::-1]
    grouped = [reversed_integer[i : i + 3] for i in range(0, len(reversed_integer), 3)]
    formatted_integer = thousands_sep.join(grouped)[::-1]
    return f"{formatted_integer}{decimal_sep}{decimal_part}"


def clean_client_name(client_name: str) -> str:
    """Cleans a raw client name for reporting."""
    if not client_name:
        return "Cliente"
    cleaned = re.sub(r"[_-][vV]\d+\b", "", client_name)
    cleaned = cleaned.replace("_", " ").replace("-", " ")
    cleaned = cleaned.strip()
    if cleaned.islower():
        return cleaned.title()
    return cleaned


def _load_asis_vocab() -> dict[str, Any]:
    p = Path(__file__).resolve().parents[4] / "engine_config" / "locales.json"
    if not p.is_file():
        p = Path(__file__).resolve().parents[4] / "engine_config" / "locales_asis_annex.json"
    if p.is_file():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            pass
    return {}


LANG_VOCAB: dict[str, Any] = _load_asis_vocab()


def set_cell_border_white(cell: Any) -> None:
    """Sets a 2.0 pt white border on all four sides of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for border_name in ["top", "left", "bottom", "right"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "16")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "FFFFFF")
        tcBorders.append(border)
    tcPr.append(tcBorders)


def set_cell_text_custom(
    cell: Any,
    text: Any,
    bold: bool = False,
    font_size: int = 9,
    align: Any = WD_ALIGN_PARAGRAPH.LEFT,
    color_rgb: Any = None,
    no_bullets: bool = False,
) -> None:
    """Populates a docx table cell with richly formatted, multi-line text."""
    clean_txt = str(text) if text is not None else ""
    lines = [line.strip() for line in clean_txt.split("\n") if line.strip()]

    is_multiline_bullet = any(
        l.startswith("- ") or l.startswith("* ") or l.startswith("• ") for l in lines
    )

    if len(lines) > 1 or (lines and is_multiline_bullet):
        cell.text = ""
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)

            clean_line = line
            bullet_prefix = ""
            if not no_bullets and (
                clean_line.startswith("- ")
                or clean_line.startswith("* ")
                or clean_line.startswith("• ")
            ):
                clean_line = clean_line[2:].strip()
                p.paragraph_format.left_indent = Inches(0.15)
                bullet_prefix = "• "
            else:
                p.paragraph_format.left_indent = Inches(0.0)
                bullet_prefix = "• "

            ref_text = ""
            if "[Ref:" in clean_line:
                clean_line, ref_part = clean_line.split("[Ref:", 1)
                ref_text = "[Ref:" + ref_part

            if ":" in clean_line:
                label, desc = clean_line.split(":", 1)
                run_lbl = p.add_run(f"{bullet_prefix}{label.strip()}: ")
                run_lbl.bold = True
                run_lbl.font.name = "Arial"
                run_lbl.font.size = Pt(font_size)
                if color_rgb:
                    run_lbl.font.color.rgb = color_rgb

                run_desc = p.add_run(desc.strip())
                run_desc.font.name = "Arial"
                run_desc.font.size = Pt(font_size)
                if color_rgb:
                    run_desc.font.color.rgb = color_rgb
            else:
                run = p.add_run(f"{bullet_prefix}{clean_line.strip()}")
                run.bold = bold
                run.font.name = "Arial"
                run.font.size = Pt(font_size)
                if color_rgb:
                    run.font.color.rgb = color_rgb

            if ref_text:
                run_ref = p.add_run(f" {ref_text.strip()}")
                run_ref.font.name = "Arial"
                run_ref.font.size = Pt(font_size - 1)
                run_ref.font.italic = True
                run_ref.font.color.rgb = RGBColor(120, 130, 140)
    else:
        set_cell_text(cell, clean_txt, bold=bold, align=align, font_size=font_size)
        p = cell.paragraphs[0]
        for r in p.runs:
            if color_rgb:
                r.font.color.rgb = color_rgb


def add_body_paragraph(
    doc: Any,
    text: Any,
    bold: bool = False,
    italic: bool = False,
    space_after: int = 6,
    text_color_rgb: tuple[int, int, int] = (46, 64, 77),
    style: str = "Normal",
) -> Any:
    """Adds a sanitized and formatted paragraph to a docx.Document object."""
    from assessment_engine.infrastructure.localization import apply_vocabulary_substitutions

    if isinstance(text, str):
        text = apply_vocabulary_substitutions(text, locale="es")

    clean_txt = str(text).replace("**", "").replace("*", "")
    is_bullet = style in ("List Bullet", "Bullet")
    target_style = "Bullet" if is_bullet else style

    try:
        p = doc.add_paragraph(style=target_style)
    except KeyError:
        p = doc.add_paragraph(style="Normal")
        if is_bullet:
            run_b = p.add_run("• ")
            run_b.font.name = "Arial"
            run_b.font.size = Pt(10)
            run_b.font.color.rgb = RGBColor(*text_color_rgb)

    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if is_bullet and ":" in clean_txt:
        parts = clean_txt.split(":", 1)
        run_bold = p.add_run(parts[0] + ":")
        run_bold.bold = True
        run_bold.font.name = "Arial"
        run_bold.font.size = Pt(10)
        run_bold.font.color.rgb = RGBColor(*text_color_rgb)

        run_rest = p.add_run(parts[1])
        run_rest.font.name = "Arial"
        run_rest.font.size = Pt(10)
        run_rest.font.color.rgb = RGBColor(*text_color_rgb)
    else:
        run = p.add_run(clean_txt)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(*text_color_rgb)
    return p


def add_callout_box(
    doc: Any,
    text: Any,
    bg_color_hex: str = "F2F2F2",
    border_color_hex: str | None = None,
    text_color_rgb: tuple[int, int, int] = (46, 64, 77),
    italic: bool = False,
    bold: bool = False,
) -> None:
    """Adds a 1x1 table representing an executive callout box with optional left border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    shade_cell(cell, bg_color_hex)

    p = cell.paragraphs[0]
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    clean_txt = str(text).replace("**", "").replace("*", "")
    is_bold = bold or ("**" in str(text))

    run = p.add_run(clean_txt)
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.bold = is_bold
    run.font.italic = italic or ("*" in str(text))
    run.font.color.rgb = RGBColor(*text_color_rgb)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")

    for b_name in ["top", "right", "bottom"]:
        b = OxmlElement(f"w:{b_name}")
        b.set(qn("w:val"), "none")
        tcBorders.append(b)

    if border_color_hex:
        left_b = OxmlElement("w:left")
        left_b.set(qn("w:val"), "single")
        left_b.set(qn("w:sz"), "36")
        left_b.set(qn("w:space"), "0")
        left_b.set(qn("w:color"), border_color_hex)
        tcBorders.append(left_b)
    else:
        left_b = OxmlElement("w:left")
        left_b.set(qn("w:val"), "none")
        tcBorders.append(left_b)

    tcPr.append(tcBorders)
    add_spacer(doc, 10)


def add_heading(
    doc: Any, text: str, level: int, primary_color_rgb: tuple[int, int, int] = (0, 114, 188)
) -> Any:
    """Adds a sanitized and styled heading to a document."""
    clean_txt = (
        text.replace("**", "").replace("__", "").replace("*", "").replace("_", "")
    )

    h = doc.add_heading(clean_txt, level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True

    if level == 1:
        h.paragraph_format.page_break_before = True

    for run in h.runs:
        run.font.color.rgb = RGBColor(*primary_color_rgb)
        run.font.name = "Arial"
        if level == 1:
            run.font.size = Pt(16)
        elif level == 2:
            run.font.size = Pt(14)
        elif level == 3:
            run.font.size = Pt(12)
        else:
            run.font.size = Pt(11)
    return h


def add_spacer(doc: Any, points: int = 12) -> Any:
    """Add a vertical spacer to a document via an empty paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(points)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    return p


def add_toc(doc: Any) -> None:
    """Inserts an Office Open XML field code to generate a Table of Contents (TOC)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run()
    fldChar = OxmlElement("w:fldChar")
    fldChar.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = 'TOC \\o "1-2" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    r_element = run._r
    r_element.append(fldChar)
    r_element.append(instrText)
    r_element.append(fldChar2)
    r_element.append(fldChar3)


def set_update_fields(doc: Any) -> None:
    """Configures the document to automatically update fields upon opening."""
    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def parse_and_append_markdown_section(
    doc: Any,
    md_path: Path,
    p_color_rgb: tuple[int, int, int],
    text_color_rgb: tuple[int, int, int],
    start_header: str | None = None,
    end_header: str | None = None,
    skip_level_1: bool = False,
) -> None:
    """Parses a specified section of a Markdown file and appends it to a DOCX document."""
    if not md_path.exists():
        return

    with open(md_path, "r", encoding="utf-8") as f:
        lines = [html.unescape(line) for line in f.readlines()]

    recording = start_header is None
    for line in lines:
        cleaned = line.strip()
        if not cleaned:
            continue

        if (
            start_header
            and cleaned.startswith("## ")
            and start_header.lower() in cleaned.lower()
        ):
            recording = True
            continue

        if (
            end_header
            and cleaned.startswith("## ")
            and end_header.lower() in cleaned.lower()
        ):
            recording = False
            break

        if not recording:
            continue

        if (
            "--- TABLA COMPARATIVA" in cleaned
            or "FORTALEZAS_CLAVE:" in cleaned
            or "BRECHAS_CLAVE:" in cleaned
            or "--- FIN TABLA" in cleaned
        ):
            continue

        if cleaned.startswith("# "):
            if skip_level_1:
                continue
            add_heading(doc, cleaned[2:], level=1, primary_color_rgb=p_color_rgb)
        elif cleaned.startswith("## "):
            add_heading(doc, cleaned[3:], level=2, primary_color_rgb=p_color_rgb)
        elif cleaned.startswith("### "):
            add_heading(doc, cleaned[4:], level=3, primary_color_rgb=p_color_rgb)
        elif cleaned.startswith("#### "):
            add_heading(doc, cleaned[5:], level=4, primary_color_rgb=p_color_rgb)
        elif cleaned.startswith("* ") or cleaned.startswith("- "):
            add_body_paragraph(
                doc, cleaned[2:], style="Bullet", text_color_rgb=text_color_rgb
            )
        elif cleaned.startswith("> "):
            text_val = cleaned[2:]
            if "Métricas de Diagnóstico Clave" in text_val:
                text_val = text_val.replace(
                    "Métricas de Diagnóstico Clave",
                    "Resumen de Exposición y Madurez Operativa",
                )
                add_callout_box(
                    doc, text_val, bg_color_hex="F4F8FA", border_color_hex="0072BC", bold=True
                )
            else:
                add_callout_box(
                    doc, text_val, bg_color_hex="F4F8FA", border_color_hex="0072BC", bold=False
                )
        elif "Bottom Line (Mensaje de Impacto):" in cleaned:
            text_val = cleaned.replace("Bottom Line (Mensaje de Impacto):", "").strip()
            add_callout_box(
                doc, text_val, bg_color_hex="EAEAEA", border_color_hex=None, bold=True
            )
        else:
            add_body_paragraph(doc, cleaned, text_color_rgb=text_color_rgb)


def add_appendix_heading(doc: Any, text: str, primary_color_rgb: Any) -> Any:
    """Adds a styled paragraph to a document, formatted as an appendix heading."""
    p = doc.add_paragraph()
    try:
        p.style = "Appendix Heading 1"
    except (OSError, ValueError, KeyError, AttributeError) as ex:
        logger.debug(f"[Fallback] Excepción controlada: {ex}")
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Century Gothic"
    run.font.size = Pt(16)
    run.font.bold = True
    if isinstance(primary_color_rgb, tuple):
        run.font.color.rgb = RGBColor(*primary_color_rgb)
    else:
        run.font.color.rgb = primary_color_rgb
    return p


def add_appendix_subheading(doc: Any, text: str, primary_color_rgb: Any) -> Any:
    """Adds a styled paragraph to a document, formatted as an appendix subheading."""
    p = doc.add_paragraph()
    try:
        p.style = "Appendix Heading 2"
    except (OSError, ValueError, KeyError, AttributeError) as ex:
        logger.debug(f"[Fallback] Excepción controlada: {ex}")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.bold = True
    if isinstance(primary_color_rgb, tuple):
        run.font.color.rgb = RGBColor(*primary_color_rgb)
    else:
        run.font.color.rgb = primary_color_rgb
    return p


def extract_key_bullets_from_md(md_path: Path, section_marker: str) -> list[str]:
    """Extracts bullet point text from a designated section of a Markdown file."""
    if not md_path.exists():
        return []
    with open(md_path, "r", encoding="utf-8") as f:
        lines = [html.unescape(line) for line in f.readlines()]
    bullets = []
    recording = False
    for line in lines:
        cleaned = line.strip()
        if section_marker in cleaned:
            recording = True
            continue
        if recording:
            if (
                cleaned.startswith("BRECHAS_CLAVE:")
                or cleaned.startswith("FORTALEZAS_CLAVE:")
                or "--- FIN TABLA" in cleaned
            ):
                break
            if cleaned.startswith("* ") or cleaned.startswith("- "):
                bullets.append(cleaned[2:])
    return bullets


def _add_markdown_bullet(
    doc: Any,
    text: str,
    text_color_rgb: tuple[int, int, int] = (46, 64, 77),
    space_after: int = 4,
) -> Any:
    """Render a bullet paragraph supporting inline bold (**) and italic (_) markdown."""
    p = doc.add_paragraph(style="Bullet")
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if ":" in text:
        label, body = text.split(":", 1)
        run_lbl = p.add_run(label + ":")
        run_lbl.bold = True
        run_lbl.font.name = "Arial"
        run_lbl.font.size = Pt(9.5)
        run_lbl.font.color.rgb = RGBColor(*text_color_rgb)
    else:
        body = text

    bold_parts = body.split("**")
    for b_idx, b_part in enumerate(bold_parts):
        is_bold = b_idx % 2 == 1
        italic_parts = b_part.split("_")
        for i_idx, i_part in enumerate(italic_parts):
            is_italic = i_idx % 2 == 1
            run = p.add_run(i_part)
            run.bold = is_bold
            run.italic = is_italic
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(*text_color_rgb)
    return p
