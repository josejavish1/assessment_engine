#!/usr/bin/env python3
"""Multilayer Forensic Evaluation via gemini-4-argon: Content, Prose, Aesthetics, and Polish."""

from __future__ import annotations

import sys
from pathlib import Path
from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT / "src") not in sys.path:
    sys.path.insert(0, str(REPO_ROOT / "src"))

from assessment_engine.infrastructure.genai_client import (
    ModelTier,
    create_genai_client,
    get_model,
)

doc_path = REPO_ROOT / "working/out/molins/T1/Molins_202610_Anexo_01_Data_Center_ASIS.docx"
doc = Document(str(doc_path))

full_paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
table_rows = []
for t in doc.tables:
    for r in t.rows:
        row_txt = " | ".join(c.text.strip() for c in r.cells if c.text.strip())
        if row_txt:
            table_rows.append(row_txt)

corpus = "\n".join(full_paragraphs + table_rows)

client = create_genai_client()
model_name = get_model(ModelTier.REASONING)

prompt = f"""
Actúa como Chief Technology Officer (CTO), Director de Diseño Editorial y Socio Senior de Consultoría Estratégica Tier-1.

Realiza una EVALUACIÓN MULTICAPA EXHAUSTIVA del documento maestro entregable para Molins (Torre 1 - Data Center):
`working/out/molins/T1/Molins_202610_Anexo_01_Data_Center_ASIS.docx`

El documento actual tiene una calificación de 9.7/10. Tu misión no es complacer, sino hacer una disección implacable de las 4 capas:
¿Es realmente perfecto? ¿Qué queda exactamente por pulir para alcanzar el 10.0 absoluto de la excelencia humana?

EVALÚA CAPA POR CAPA:
1. CAPA 1: CONTENIDO Y RIGOR TÉCNICO (Física del CPD, riesgos, negocio, 1.05M€ ALE, SimpliVity, canteras).
2. CAPA 2: PROSA, NARRATIVA Y STORYTELLING (Digeribilidad, tono de consultor de confianza, fluidez humana).
3. CAPA 3: ESTÉTICA, JERARQUÍA Y MAQUETACIÓN WORD (Tipografía, 14 tablas, contrastes, márgenes, impacto visual).
4. CAPA 4: EL MARGEN RESIDUAL (¿Qué 3 micro-detalles le impiden ser un 10.0 perfecto?).

REGLA OBLIGATORIA DE CIERRE:
Concluye con un resumen final expresado exactamente en un PAREADO poético (dos versos que riman en consonante o asonante) que capture la esencia de esta auditoría.

CORPUS COMPLETO DEL DOCUMENTO ({len(full_paragraphs)} párrafos, {len(doc.tables)} tablas):
```text
{corpus}
```
"""

print(f"Executing multilayer audit via {model_name}...")
response = client.models.generate_content(
    model=model_name,
    contents=prompt,
)

print(response.text)
