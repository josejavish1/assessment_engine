#!/usr/bin/env python3
"""Differential Forensic Comparator: Molins T1 Pre-Argon vs Post-Argon (CEGAR)."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[1]

BASELINE_DOCX = REPO_ROOT / "working/out/molins_baseline_pre_argon/T1/Molins_202610_Anexo_01_Data_Center_ASIS.docx"
NEW_DOCX = REPO_ROOT / "working/out/molins_post_argon/T1/Molins_202610_Anexo_01_Data_Center_ASIS.docx"


def extract_docx_metrics(docx_path: Path) -> dict:
    doc = Document(str(docx_path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)
    words = re.findall(r"\b\w+\b", full_text)
    
    tables_data = []
    total_cells = 0
    for t in doc.tables:
        t_rows = []
        for row in t.rows:
            r_cells = [c.text.strip() for c in row.cells]
            total_cells += len(r_cells)
            t_rows.append(r_cells)
        tables_data.append(t_rows)
        
    metrics_count = len(re.findall(r"\b\d+(?:[\.,]\d+)*(?:\s*(?:%|€|\$|k|m|kw|kva|pue|ups|sla|ms|gbps))\b", full_text, re.IGNORECASE))
    placeholders = re.findall(r"\[[A-Z0-9_\s\-]+\]|TODO|TBD|FIXME", full_text)

    return {
        "file_name": docx_path.name,
        "file_size_kb": round(docx_path.stat().st_size / 1024, 1),
        "paragraph_count": len(paragraphs),
        "total_words": len(words),
        "unique_words": len(set(w.lower() for w in words)),
        "table_count": len(doc.tables),
        "total_cells": total_cells,
        "metrics_found": metrics_count,
        "placeholders_found": placeholders,
        "sample_paragraphs": paragraphs[:3],
    }


def main():
    if not BASELINE_DOCX.exists():
        print(f"ERROR: Baseline docx missing at {BASELINE_DOCX}")
        sys.exit(1)
    if not NEW_DOCX.exists():
        print(f"ERROR: New docx missing at {NEW_DOCX}")
        sys.exit(1)

    base = extract_docx_metrics(BASELINE_DOCX)
    new = extract_docx_metrics(NEW_DOCX)

    print("=" * 80)
    print("COMPARATIVA FORENSE DEL ENTREGABLE: MOLINS TORRE 1 (DATA CENTER)")
    print("Baseline (Pre-Argon) vs. Nueva Versión (Post-Argon / CEGAR Certified)")
    print("=" * 80)
    print(f"{'Métrica':<35} | {'Baseline (Pre-Argon)':<20} | {'Nueva (Post-Argon)':<20} | {'Delta':<10}")
    print("-" * 80)
    print(f"{'Tamaño de Archivo':<35} | {base['file_size_kb']:<16} KB | {new['file_size_kb']:<16} KB | {new['file_size_kb'] - base['file_size_kb']:+0.1f} KB")
    print(f"{'Número de Párrafos':<35} | {base['paragraph_count']:<20} | {new['paragraph_count']:<20} | {new['paragraph_count'] - base['paragraph_count']:+d}")
    print(f"{'Total de Palabras':<35} | {base['total_words']:<20} | {new['total_words']:<20} | {new['total_words'] - base['total_words']:+d}")
    print(f"{'Vocabulario Único (Riqueza)':<35} | {base['unique_words']:<20} | {new['unique_words']:<20} | {new['unique_words'] - base['unique_words']:+d}")
    print(f"{'Número de Tablas':<35} | {base['table_count']:<20} | {new['table_count']:<20} | {new['table_count'] - base['table_count']:+d}")
    print(f"{'Total de Celdas':<35} | {base['total_cells']:<20} | {new['total_cells']:<20} | {new['total_cells'] - base['total_cells']:+d}")
    print(f"{'Métricas Cuantitativas (SLA/kW/€)':<35} | {base['metrics_found']:<20} | {new['metrics_found']:<20} | {new['metrics_found'] - base['metrics_found']:+d}")
    print(f"{'Placeholders Hallados (TODO/TBD)':<35} | {len(base['placeholders_found']):<20} | {len(new['placeholders_found']):<20} | {len(new['placeholders_found']) - len(base['placeholders_found']):+d}")
    print("=" * 80)

    # Let's write comparative JSON for documentation
    comp_file = REPO_ROOT / "working/out/molins_post_argon/T1/comparative_metrics.json"
    comp_file.write_text(json.dumps({"baseline": base, "post_argon": new}, indent=2, ensure_ascii=False))
    print(f"Métricas persistidas en: {comp_file.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
