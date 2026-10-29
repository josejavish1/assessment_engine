"""Empirical simulation harness for multinational, multi-sector, and multi-lingual archetypes.

Executes end-to-end artifact generation and physical DOCX validation for 4 global enterprise archetypes:
1. Nordic Grid Power (Germany / Critical Infrastructure / de / EUR / NIS2)
2. Global Nippon Finance (Japan / Banking / ja / JPY / Basel III)
3. Meridian BioHealth (USA / Pharma / en / USD / HIPAA)
4. Iberia Retail Logistics (Spain / Retail / es / EUR / ENS Alta)
"""

from __future__ import annotations

import json
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
from docx import Document

from assessment_engine.adapters.compilers.docx_compiler import DocxCompiler
from assessment_engine.adapters.render_commercial_report import render_commercial_report
from assessment_engine.adapters.render_tower_annex_from_template import (
    main as render_annex_main,
)
from assessment_engine.domain.schemas.ast import (
    CellNode,
    DocumentAST,
    HeadingNode,
    ParagraphNode,
    SpacerNode,
    TableNode,
    TableRowNode,
)
from assessment_engine.domain.schemas.commercial import (
    CommercialDocumentMeta,
    CommercialPayload,
    CommercialSummaryDraft,
    ContextAndWhy,
    DealFlash,
    DeliveryTeam,
    GovernanceAndAssumptions,
    GtmStrategy,
    InvestmentAndTimeline,
    OpportunityPipelineItem,
    ProposalDraft,
    RiskManagementItem,
    ScopeAndHow,
    SolutionAndWhat,
    StakeholderMatrixItem,
    WhyPartner,
)
from assessment_engine.infrastructure.policies.fair_risk import FairRiskPolicy
from assessment_engine.infrastructure.text_utils import format_currency_custom, slugify

ROOT = Path(__file__).resolve().parent.parent
BASE_OUTPUT_DIR = ROOT / "working" / "out" / "simulations"
ANNEX_TEMPLATE_PATH = ROOT / "templates" / "docx" / "tower_annex_template.docx"
COMMERCIAL_TEMPLATE_PATH = ROOT / "templates" / "docx" / "global_report_template.docx"


@dataclass
class ArchetypeConfig:
    """Configuration parameters for a simulated multinational enterprise archetype."""

    archetype_id: str
    company_name: str
    country: str
    jurisdiction: str
    industry: str
    doc_lang: str
    currency_symbol: str
    currency_code: str
    regulatory_frameworks: list[str]
    critical_systems: list[str]
    strategic_objectives: list[str]
    hourly_rate_sre: float
    estimated_tam_str: str
    estimated_capex_str: str
    estimated_opex_reduction_str: str
    tcv_range_str: str


ARCHETYPES: list[ArchetypeConfig] = [
    ArchetypeConfig(
        archetype_id="ARCH-DE-GRID",
        company_name="Nordic Grid Power GmbH",
        country="Germany",
        jurisdiction="European Union (DE)",
        industry="critical_infrastructure",
        doc_lang="de",
        currency_symbol="€",
        currency_code="EUR",
        regulatory_frameworks=["NIS2", "ISO 27001", "BSI IT-Grundschutz"],
        critical_systems=["SCADA Energy Dispatcher", "Siemens Spectrum Power", "Substation RTUs"],
        strategic_objectives=[
            "Härtung von OT-Netzwerken nach IEC 62443",
            "Automatisierung der Netzbetriebsführung",
            "Resilienz gegen koordinierte Cyberangriffe",
        ],
        hourly_rate_sre=48.0,
        estimated_tam_str="2,5M €",
        estimated_capex_str="1,2M €",
        estimated_opex_reduction_str="350K € / Jahr",
        tcv_range_str="850.000 € - 1.100.000 €",
    ),
    ArchetypeConfig(
        archetype_id="ARCH-JP-FIN",
        company_name="Global Nippon Financial Group",
        country="Japan",
        jurisdiction="Japan / APAC",
        industry="banking",
        doc_lang="ja",
        currency_symbol="¥",
        currency_code="JPY",
        regulatory_frameworks=["FISC Security Guidelines", "Basel III", "ISO 27001"],
        critical_systems=["Mainframe Core Settlement", "ZEDI Foreign Exchange", "Swift Gateway"],
        strategic_objectives=[
            "勘定系システムのハイブリッドクラウド移行",
            "API基盤のマイクロサービス化と耐障害性向上",
            "リアルタイム取引モニタリングと不正検知の自動化",
        ],
        hourly_rate_sre=7500.0,
        estimated_tam_str="350M ¥",
        estimated_capex_str="180M ¥",
        estimated_opex_reduction_str="45M ¥ / 年",
        tcv_range_str="120,000,000¥ - 160,000,000¥",
    ),
    ArchetypeConfig(
        archetype_id="ARCH-US-HEALTH",
        company_name="Meridian BioHealth Inc.",
        country="United States",
        jurisdiction="North America (FDA / HHS)",
        industry="healthcare_pharma",
        doc_lang="en",
        currency_symbol="$",
        currency_code="USD",
        regulatory_frameworks=["HIPAA", "FDA 21 CFR Part 11", "NIST CSF"],
        critical_systems=["Epic EHR Core", "Veeva Clinical Vault", "Genomic Sequencing HPC"],
        strategic_objectives=[
            "Zero Trust Segmentation across Clinical Perimeters",
            "HIPAA Audit Trail Automation in Hybrid Cloud",
            "Accelerate Drug Discovery Pipeline Resilience",
        ],
        hourly_rate_sre=55.0,
        estimated_tam_str="$2.5M",
        estimated_capex_str="$1.2M",
        estimated_opex_reduction_str="$350K / year",
        tcv_range_str="$850,000 - $1,100,000",
    ),
    ArchetypeConfig(
        archetype_id="ARCH-ES-RETAIL",
        company_name="Iberia Retail Logistics S.A.",
        country="Spain",
        jurisdiction="European Union (ES)",
        industry="retail",
        doc_lang="es",
        currency_symbol="€",
        currency_code="EUR",
        regulatory_frameworks=["ENS Alta", "GDPR", "PCI-DSS"],
        critical_systems=["SAP S/4HANA Supply Chain", "Manhattan WMS", "Terminales POS en Tiendas"],
        strategic_objectives=[
            "Garantizar disponibilidad 99.99% en campaña de Black Friday",
            "Consolidar inventario omnicanal en tiempo real",
            "Cumplimiento riguroso del Esquema Nacional de Seguridad",
        ],
        hourly_rate_sre=38.0,
        estimated_tam_str="2,5M €",
        estimated_capex_str="1,2M €",
        estimated_opex_reduction_str="350K € / año",
        tcv_range_str="850.000 € - 1.100.000 €",
    ),
]


def extract_docx_text(path: Path) -> str:
    """Extract all text runs from paragraphs and tables in a DOCX document."""
    doc = Document(str(path))
    paras = [p.text for p in doc.paragraphs]
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(paras + cells)


def simulate_archetype(cfg: ArchetypeConfig) -> dict[str, Any]:
    """Generates physical deliverables for an archetype and runs automated validation."""
    slug = slugify(cfg.company_name)
    out_dir = BASE_OUTPUT_DIR / slug
    out_dir.mkdir(parents=True, exist_ok=True)

    audit_results: dict[str, Any] = {
        "archetype_id": cfg.archetype_id,
        "company_name": cfg.company_name,
        "language": cfg.doc_lang,
        "currency": cfg.currency_symbol,
        "generated_files": [],
        "checks_passed": True,
        "check_details": [],
    }

    # 1. FAIR Actuarial Risk Simulation
    fair_policy = FairRiskPolicy()
    rng = np.random.default_rng(42)
    tef = fair_policy._sample_pert(2.0, 5.0, 12.0, size=5000, rng=rng)
    vuln = fair_policy._sample_pert(0.15, 0.30, 0.60, size=5000, rng=rng)
    loss = fair_policy._sample_lognormal(
        rng,
        cfg.hourly_rate_sre * 200,
        cfg.hourly_rate_sre * 800,
        cfg.hourly_rate_sre * 3500,
        size=5000,
    )
    losses = tef * vuln * loss
    ale_mean = float(np.mean(losses))
    p95_loss = float(np.percentile(losses, 95))

    ale_formatted = format_currency_custom(
        ale_mean,
        {"currency_symbol": cfg.currency_symbol},
        doc_lang=cfg.doc_lang,
    )

    fair_data = {
        "ale_mean": ale_mean,
        "ale_formatted": ale_formatted,
        "p95_loss": p95_loss,
        "iterations": 5000,
        "currency": cfg.currency_code,
    }
    fair_json_path = out_dir / "fair_actuarial_model.json"
    fair_json_path.write_text(json.dumps(fair_data, indent=2, ensure_ascii=False), encoding="utf-8")
    audit_results["generated_files"].append(str(fair_json_path))

    # 2. Build Commercial Proposal Deliverable
    comm_payload = CommercialPayload(
        meta=CommercialDocumentMeta(
            client=cfg.company_name,
            date="2026-10-28",
            version="1.0.0",
            doc_lang=cfg.doc_lang,
            currency_symbol=cfg.currency_symbol,
        ),
        commercial_summary=CommercialSummaryDraft(
            deal_flash=DealFlash(
                purchase_driver=f"Transformación de {cfg.company_name} bajo los estándares {', '.join(cfg.regulatory_frameworks)}.",
                vendor_win_theme="Liderazgo en Automatización Cloud y Soberanía Tecnológica",
            ),
            why_now_bullets=cfg.strategic_objectives,
            how_we_win_bullets=[
                f"Capacidad comprobada en {cfg.industry}",
                "Aceleradores de observabilidad y soberanía tecnológica",
                "Cumplimiento automatizado de normativas",
            ],
            estimated_tam=cfg.estimated_tam_str,
        ),
        gtm_strategy=GtmStrategy(
            trojan_horse="Piloto controlado en infraestructura crítica",
            self_funded_transformation=f"Optimización de costes para autofinanciar la fase 2 ({ale_formatted} mitigados)",
            lock_in=f"Desacoplamiento progresivo mediante APIs abiertas para soberanía en {cfg.jurisdiction}",
        ),
        stakeholder_matrix=[
            StakeholderMatrixItem(
                name="CIO / CISO",
                role="Sponsor Ejecutivo",
                focus=f"Cumplimiento {cfg.regulatory_frameworks[0]} y Disponibilidad",
                message="Propuesta de valor verificada empíricamente",
            )
        ],
        opportunities_pipeline=[
            OpportunityPipelineItem(
                initiative=f"Modernización de {cfg.critical_systems[0]}",
                tower_origin="T1",
                vendor_cosell="Enterprise Architecture Core",
                revenue_type="Servicios Profesionales",
                estimated_tcv=cfg.estimated_capex_str,
                cost_of_inaction=f"ALE proyectado de {ale_formatted}",
                client_objection="Riesgo de indisponibilidad en producción",
                objection_handling="Despliegue Canary y rollback automatizado",
            )
        ],
        proactive_proposals=[
            ProposalDraft(
                initiative_name=f"Iniciativa Núcleo para {cfg.critical_systems[0]}",
                context_and_why=ContextAndWhy(
                    origin=f"Auditoría técnica de {cfg.industry}",
                    cost_of_inaction=f"Exposición a pérdidas de {ale_formatted}",
                ),
                solution_and_what=SolutionAndWhat(
                    target_state=f"Arquitectura distribuida resiliente para {cfg.critical_systems[0]}.",
                    north_star_metric="Disponibilidad 99.99% y RTO < 15 min",
                ),
                scope_and_how=ScopeAndHow(
                    phases=["Fase 1: Diagnóstico", "Fase 2: Piloto", "Fase 3: Producción"],
                    deliverables=["Documento HLD/LLD", "Plan de contingencia"],
                    out_of_scope=["Desarrollo de nuevas aplicaciones de usuario"],
                ),
                delivery_team=DeliveryTeam(team_roles=["Principal Enterprise Architect", "Lead SRE Engineer"]),
                ai_transformation_strategy="Aceleración con agentes supervisores y Gemini Code Assist.",
                governance_and_assumptions=GovernanceAndAssumptions(
                    governance_model="Comité de seguimiento quincenal",
                    assumptions=["Acceso a especificaciones de sistemas críticos"],
                ),
                risk_management=[
                    RiskManagementItem(
                        risk="Interrupción durante la migración",
                        mitigation="Estrategia Canary Deployment con rollback instantáneo.",
                    )
                ],
                activation_plan=["Fase de Diagnóstico", "Implementación Piloto", "Pase a Producción"],
                why_ntt_data=WhyPartner(
                    accelerators=["Framework de automatización PaC / IaC"],
                    partnerships="Alianza estratégica global Tier-1 con Hyperscalers",
                ),
                investment_and_timeline=InvestmentAndTimeline(
                    estimated_duration="6 meses",
                    tcv_range=cfg.tcv_range_str,
                ),
                executive_synthesis=f"Retorno demostrable y mitigación directa del riesgo de {ale_formatted} en el primer año.",
            )
        ],
        intelligence_dossier={"systems": cfg.critical_systems, "jurisdiction": cfg.jurisdiction},
    )

    comm_docx_path = out_dir / f"Propuesta_Comercial_{slug}.docx"
    render_commercial_report(comm_payload, COMMERCIAL_TEMPLATE_PATH, comm_docx_path)
    audit_results["generated_files"].append(str(comm_docx_path))

    # 3. Build Technical Annex Deliverable (via template renderer)
    annex_payload_data = json.loads(Path("tests/test_data/approved_annex_t2.template_payload.json").read_text(encoding="utf-8"))
    annex_payload_data["document_meta"]["client_name"] = cfg.company_name
    annex_payload_data["document_meta"]["language"] = cfg.doc_lang
    annex_payload_data["document_meta"]["regulatory_context"] = cfg.regulatory_frameworks
    annex_payload_data["document_meta"]["source_client"] = "redeia_v3"

    annex_payload_str = json.dumps(annex_payload_data, ensure_ascii=False)
    for src_name in ["redeia_v3", "REDEIA", "Redeia", "Red Eléctrica"]:
        annex_payload_str = annex_payload_str.replace(src_name, cfg.company_name)
    if cfg.currency_symbol != "€":
        annex_payload_str = annex_payload_str.replace("€", cfg.currency_symbol)
    annex_payload_data = json.loads(annex_payload_str)

    annex_json_path = out_dir / "approved_annex_payload.json"
    annex_json_path.write_text(json.dumps(annex_payload_data, indent=2, ensure_ascii=False), encoding="utf-8")

    annex_docx_path = out_dir / f"Anexo_Tecnico_{slug}.docx"
    render_annex_main(["render_tower_annex", str(annex_json_path), str(ANNEX_TEMPLATE_PATH), str(annex_docx_path)])
    audit_results["generated_files"].append(str(annex_docx_path))

    # 4. Build AST Technical Blueprint Deliverable
    ast = DocumentAST(
        title=f"Technical Blueprint & Enterprise Architecture: {cfg.company_name}",
        metadata={
            "client": cfg.company_name,
            "version": "1.0",
            "language": cfg.doc_lang,
            "currency": cfg.currency_symbol,
            "industry": cfg.industry,
        },
        nodes=[
            HeadingNode(text=f"1. Executive Architecture Synthesis - {cfg.company_name}", level=1),
            ParagraphNode(
                text=(
                    f"Die Organisation {cfg.company_name} betreibt kritische Systeme im Sektor {cfg.industry}. "
                    f"Die quantifizierte jährliche Verlustschätzung (ALE) beläuft sich auf {ale_formatted}. "
                    f"Die Einhaltung von {', '.join(cfg.regulatory_frameworks)} ist zwingend erforderlich."
                    if cfg.doc_lang == "de"
                    else (
                        f"{cfg.company_name} operates critical systems in the {cfg.industry} sector. "
                        f"The quantified Annualized Loss Expectancy (ALE) stands at {ale_formatted}. "
                        f"Compliance with {', '.join(cfg.regulatory_frameworks)} is mandatory."
                        if cfg.doc_lang == "en"
                        else (
                            f"{cfg.company_name}は{cfg.industry}分野における重要基幹システムを運用しています。"
                            f"定量化された年間損失予測（ALE）は{ale_formatted}です。"
                            f"{', '.join(cfg.regulatory_frameworks)}への厳格な準拠が必須要件となります。"
                            if cfg.doc_lang == "ja"
                            else f"La organización {cfg.company_name} opera infraestructura crítica en el sector {cfg.industry}. La pérdida anual esperada (ALE) asciende a {ale_formatted}. El cumplimiento de {', '.join(cfg.regulatory_frameworks)} es de máxima prioridad."
                        )
                    )
                ),
                style="Normal",
                bold=False,
            ),
            SpacerNode(points=12),
            HeadingNode(text="2. Critical Target Systems & Invariants", level=2),
            TableNode(
                rows=[
                    TableRowNode(
                        cells=[
                            CellNode(text="Critical System", bold=True, bg_color="D9EAF7", align="CENTER"),
                            CellNode(text="Target Objective", bold=True, bg_color="D9EAF7", align="CENTER"),
                            CellNode(text="Regulatory Shield", bold=True, bg_color="D9EAF7", align="CENTER"),
                        ],
                        is_header=True,
                    ),
                    *[
                        TableRowNode(
                            cells=[
                                CellNode(text=sys_name, bold=True),
                                CellNode(text=f"High Availability & Resilient SRE for {sys_name}"),
                                CellNode(text=cfg.regulatory_frameworks[i % len(cfg.regulatory_frameworks)]),
                            ]
                        )
                        for i, sys_name in enumerate(cfg.critical_systems)
                    ],
                ]
            ),
        ],
    )
    ast_compiler = DocxCompiler()
    blueprint_docx_path = out_dir / f"Technical_Blueprint_{slug}.docx"
    ast_compiler.compile(ast, str(blueprint_docx_path))
    audit_results["generated_files"].append(str(blueprint_docx_path))

    # --- AUTOMATED EMPIRICAL AUDITS ---
    docx_files = [comm_docx_path, annex_docx_path, blueprint_docx_path]

    for d_path in docx_files:
        assert d_path.is_file(), f"Physical file missing: {d_path}"
        f_size = d_path.stat().st_size
        assert f_size > 10000, f"File {d_path.name} is abnormally small ({f_size} bytes)"

        extracted_text = extract_docx_text(d_path)

        # Audit 1: Zero Template Leaks
        has_leak = "{{" in extracted_text or "}}" in extracted_text or "{client_name}" in extracted_text
        audit_results["check_details"].append({
            "file": d_path.name,
            "check": "Zero Template Leaks",
            "passed": not has_leak,
        })
        if has_leak:
            audit_results["checks_passed"] = False

        # Audit 2: Client Name Integrity
        has_client = cfg.company_name.upper() in extracted_text.upper() or slug in extracted_text.lower()
        audit_results["check_details"].append({
            "file": d_path.name,
            "check": "Client Name Integrity",
            "passed": has_client,
        })
        if not has_client:
            audit_results["checks_passed"] = False

        # Audit 3: Currency Symbol Presence
        has_curr = cfg.currency_symbol in extracted_text
        audit_results["check_details"].append({
            "file": d_path.name,
            "check": "Currency Symbol Invariance",
            "passed": has_curr,
        })
        if not has_curr:
            audit_results["checks_passed"] = False

        # Audit 4: Anti-Buzzword Sanitation
        has_buzzword = "sinergias fabriles" in extracted_text or "palanca de valor" in extracted_text
        audit_results["check_details"].append({
            "file": d_path.name,
            "check": "Anti-Buzzword Sanitation",
            "passed": not has_buzzword,
        })
        if has_buzzword:
            audit_results["checks_passed"] = False

        # Audit 5: Zero Residual Client Leakage (purgar clientes previos como Redeia o Secil)
        has_client_leak = "redeia" in extracted_text.lower() or "secil" in extracted_text.lower()
        audit_results["check_details"].append({
            "file": d_path.name,
            "check": "Zero Residual Client Leakage",
            "passed": not has_client_leak,
        })
        if has_client_leak:
            audit_results["checks_passed"] = False

        # Audit 6: Multilingual Cover Page Integrity for Commercial Proposals
        if d_path == comm_docx_path:
            if cfg.doc_lang == "en":
                passed_cov = "Classification: CONFIDENTIAL - INTERNAL USE ONLY" in extracted_text and "Fecha:" not in extracted_text
            elif cfg.doc_lang == "de":
                passed_cov = "Klassifizierung: VERTRAULICH - NUR FÜR INTERNEN GEBRAUCH" in extracted_text and "Fecha:" not in extracted_text
            elif cfg.doc_lang == "ja":
                passed_cov = "機密区分: 機密 - 社内限定" in extracted_text and "Fecha:" not in extracted_text
            else:
                passed_cov = "Clasificación: CONFIDENCIAL - USO INTERNO" in extracted_text
            audit_results["check_details"].append({
                "file": d_path.name,
                "check": "Multilingual Cover Page Integrity",
                "passed": passed_cov,
            })
            if not passed_cov:
                audit_results["checks_passed"] = False

    return audit_results


def main() -> None:
    print("=" * 80)
    print("  MULTINATIONAL ARCHETYPE EMPIRICAL SIMULATION HARNESS")
    print("  Countries: DE, JP, US, ES | Sectors: Critical Infra, Banking, Pharma, Retail")
    print("=" * 80)

    BASE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    summary_results = []
    all_passed = True

    for archetype in ARCHETYPES:
        print(f"\n🚀 Simulating {archetype.company_name} ({archetype.country} / {archetype.doc_lang})...")
        res = simulate_archetype(archetype)
        summary_results.append(res)
        if not res["checks_passed"]:
            all_passed = False
            print(f"❌ Archetype {archetype.archetype_id} failed quality checks!")
        else:
            print(f"✅ Generated {len(res['generated_files'])} artifacts. All empirical checks passed.")

    summary_file = BASE_OUTPUT_DIR / "SIMULATION_SUMMARY.json"
    existing_entries = []
    if summary_file.exists():
        try:
            existing_entries = json.loads(summary_file.read_text(encoding="utf-8"))
        except Exception:
            existing_entries = []

    canonical_ids = {a.archetype_id for a in ARCHETYPES}
    merged_summary = [e for e in existing_entries if e.get("archetype_id") not in canonical_ids] + summary_results
    summary_file.write_text(json.dumps(merged_summary, indent=2, ensure_ascii=False), encoding="utf-8")

    print("\n" + "=" * 80)
    print(f"  SIMULATION SUMMARY: {'ALL 4 ARCHETYPES CERTIFIED' if all_passed else 'FAILURES DETECTED'}")
    print(f"  Deliverables isolated at: {BASE_OUTPUT_DIR}")
    print(f"  Audit log: {summary_file}")
    print("=" * 80)

    if not all_passed:
        sys.exit(1)


if __name__ == "__main__":
    main()
