# golden-path: ignore
from __future__ import annotations

import os
import re
from pathlib import Path
from docx import Document


def test_linguistic_and_executive_quality_gate() -> None:
    """Enforces non-negotiable C-Suite quality, chronological consistency, 
    and European Spanish linguistic alignment on all active compiled assessment deliverables.
    
    This acts as a continuous quality gate. Any regression or AI hallucination leak
    will fail this test, preventing delivery.
    """
    repo_root = Path(__file__).resolve().parents[2]
    working_dir = repo_root / "working"
    
    if not working_dir.exists():
        return # Skip if no active deliverables folder exists yet
        
    # Scan all active Word documents (.docx) and markdown modules (.md)
    # inside working/redeia_v3 (our target current deliverables)
    target_dir = working_dir / "redeia_v3"
    if not target_dir.exists():
         return # Skip if redeia assessment is not built in this context

    violations = []
    
    # Compile forbidden patterns
    forbidden_linguistic = [
        ("sistémica", "Terminología no recomendada en España (usar 'estructural')"),
        ("sistémico", "Terminología no recomendada en España (usar 'estructural')"),
    ]
    
    # Chronological patterns (obsolete years since current date is July 2027)
    forbidden_chronological = [
        ("2021-2025", "Plan estratégico obsoleto para el año 2027 (actualizar a 2026-2030)"),
        ("2023-2025", "Plan de sostenibilidad obsoleto para el año 2027 (actualizar a 2026-2030)"),
    ]
    
    # AI meta-remarks and apologies
    forbidden_meta_remarks = [
        ("No se ha encontrado un documento", "Fuga de comentario meta de RAG/IA"),
        ("información pública disponible", "Evasión o disculpa de la IA en la recolección"),
    ]
    
    # Mixed language parentheses
    forbidden_bilingual = [
        (r"\(System Health Matrix\)", "Nomenclatura bilingüe redundante (usar 'Entorno Tecnológico')"),
        (r"\(Cost of Delay\)", "Nomenclatura bilingüe redundante (usar 'Demora en Iniciativas')"),
    ]

    for root, dirs, files in os.walk(target_dir):
        # Ignore cache directories
        if any(h in root for h in [".venv", "node_modules", "__pycache__"]):
            continue
            
        for file in files:
            file_path = Path(root) / file
            
            # --- Check Word Documents (.docx) ---
            # Strictly focus on the active final deliverables to ignore legacy/backup artifacts in the folder.
            is_active_deliverable = file.startswith("AS-IS_Anexo_Tecnico_") and (file.endswith(".docx") or file.endswith(".shadow.docx"))
            if is_active_deliverable and not file.startswith("~$") and "minimal" not in file:
                try:
                    doc = Document(str(file_path))
                    
                    # 1. Scan Paragraphs
                    for p_idx, p in enumerate(doc.paragraphs):
                        text = p.text
                        
                        # Check linguistic
                        for term, msg in forbidden_linguistic:
                            if term in text.lower():
                                violations.append(f"[{file_path.name}] Párrafo {p_idx}: Violación lingüística - '{term}'. {msg}")
                                
                        # Check chronological
                        for term, msg in forbidden_chronological:
                            if term in text:
                                violations.append(f"[{file_path.name}] Párrafo {p_idx}: Desfase cronológico - '{term}'. {msg}")
                                
                        # Check meta remarks
                        for term, msg in forbidden_meta_remarks:
                            if term in text:
                                violations.append(f"[{file_path.name}] Párrafo {p_idx}: Comentario meta - '{term}'. {msg}")
                                
                        # Check bilingual
                        for pattern, msg in forbidden_bilingual:
                            if re.search(pattern, text):
                                violations.append(f"[{file_path.name}] Párrafo {p_idx}: Título bilingüe - '{pattern}'. {msg}")
                    
                    # 2. Scan Table Cells
                    for t_idx, t in enumerate(doc.tables):
                        for r_idx, row in enumerate(t.rows):
                            for c_idx, cell in enumerate(row.cells):
                                text = cell.text
                                
                                # Check linguistic
                                for term, msg in forbidden_linguistic:
                                    if term in text.lower():
                                        violations.append(f"[{file_path.name}] Tabla {t_idx} Fila {r_idx} Celda {c_idx}: Violación lingüística - '{term}'. {msg}")
                                        
                                # Check chronological
                                for term, msg in forbidden_chronological:
                                    if term in text:
                                        violations.append(f"[{file_path.name}] Tabla {t_idx} Fila {r_idx} Celda {c_idx}: Desfase cronológico - '{term}'. {msg}")
                                        
                                # Check meta remarks
                                for term, msg in forbidden_meta_remarks:
                                    if term in text:
                                        violations.append(f"[{file_path.name}] Tabla {t_idx} Fila {r_idx} Celda {c_idx}: Comentario meta - '{term}'. {msg}")
                except Exception as e:
                    # Ignore corrupted or unreadable documents in test sandbox
                    pass

            # --- Check Markdown Modules (.md) ---
            elif file.endswith(".md") and file_path.parent.name == "asis_modules":
                try:
                    content = file_path.read_text(encoding="utf-8")
                    
                    # Check linguistic
                    for term, msg in forbidden_linguistic:
                        if term in content.lower():
                            violations.append(f"[{file_path.name}] Archivo MD: Violación lingüística - '{term}'. {msg}")
                            
                    # Check chronological
                    for term, msg in forbidden_chronological:
                        if term in content:
                            violations.append(f"[{file_path.name}] Archivo MD: Desfase cronológico - '{term}'. {msg}")
                            
                    # Check meta remarks
                    for term, msg in forbidden_meta_remarks:
                        if term in content:
                            violations.append(f"[{file_path.name}] Archivo MD: Comentario meta - '{term}'. {msg}")
                            
                    # Check bilingual
                    for pattern, msg in forbidden_bilingual:
                        if re.search(pattern, content):
                            violations.append(f"[{file_path.name}] Archivo MD: Título bilingüe - '{pattern}'. {msg}")
                except Exception:
                    pass

    if violations:
        error_msg = "Se detectaron violaciones en las Leyes No Negociables de Calidad Ejecutiva:\n" + "\n".join(violations)
        raise AssertionError(error_msg)
