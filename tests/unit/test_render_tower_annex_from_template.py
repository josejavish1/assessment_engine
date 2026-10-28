"""Unit test for rendering tower annex from template."""

from __future__ import annotations

from pathlib import Path
from docx import Document

from assessment_engine.adapters.render_tower_annex_from_template import main


def test_render_tower_annex_from_template_execution(tmp_path: Path):
    """Verify that main() renders a tower annex DOCX document from payload and template."""
    payload_path = Path("tests/test_data/approved_annex_t2.template_payload.json")
    template_path = Path("templates/docx/tower_annex_template.docx")
    output_docx = tmp_path / "Annex_T2_rendered.docx"

    assert payload_path.exists(), "Payload fixture missing."
    assert template_path.exists(), "Template DOCX missing."

    main(["render_tower_annex", str(payload_path), str(template_path), str(output_docx)])

    assert output_docx.is_file(), "El archivo DOCX del anexo debe haber sido generado físicamente."

    doc = Document(str(output_docx))
    full_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        cell.text for t in doc.tables for r in t.rows for cell in r.cells
    )

    # Invariantes semánticos de la torre
    assert "REDEIA" in full_text.upper() or "RED ELÉCTRICA" in full_text.upper()
    assert "T2" in full_text
    assert "Definido" in full_text
    assert "Gestionado" in full_text
    assert "ALE" in full_text

    # Invariante de Cero Fugas de Plantilla
    assert "{{" not in full_text, "Variables de plantilla {{...}} no resueltas en el anexo de torre."
    assert "}}" not in full_text, "Variables de plantilla {{...}} no resueltas en el anexo de torre."
