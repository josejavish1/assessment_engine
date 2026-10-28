import docx
import pytest

from assessment_engine.infrastructure.docx_render_utils import (
    render_gap_table,
    render_pillar_score_table,
)


def test_render_pillar_score_table_bilingual_headers():
    """Verify that render_pillar_score_table uses localized headers based on locale."""
    # 1. English headers
    doc_en = docx.Document()
    p_en = doc_en.add_paragraph("{{TABLA_PILARES}}")
    render_pillar_score_table(doc_en, "{{TABLA_PILARES}}", [], locale="en")

    table_en = doc_en.tables[0]
    headers_en = [cell.text for cell in table_en.rows[0].cells]
    assert headers_en == ["Pillar", "Score", "Level", "Executive Summary"]

    # 2. Spanish headers
    doc_es = docx.Document()
    p_es = doc_es.add_paragraph("{{TABLA_PILARES}}")
    render_pillar_score_table(doc_es, "{{TABLA_PILARES}}", [], locale="es")

    table_es = doc_es.tables[0]
    headers_es = [cell.text for cell in table_es.rows[0].cells]
    assert headers_es == ["Pilar", "Score", "Nivel", "Lectura ejecutiva"]


def test_render_gap_table_bilingual_headers():
    """Verify that render_gap_table uses localized headers based on locale."""
    # English gap headers
    doc_en = docx.Document()
    p_en = doc_en.add_paragraph("{{TABLA_BRECHAS}}")
    render_gap_table(doc_en, "{{TABLA_BRECHAS}}", [], locale="en")

    table_en = doc_en.tables[0]
    headers_en = [cell.text for cell in table_en.rows[0].cells]
    assert headers_en == ["Pillar", "Current State (AS-IS)", "Target State (TO-BE)", "Key Gap"]
