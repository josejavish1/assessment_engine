from __future__ import annotations

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from assessment_engine.adapters.compilers.docx_compiler import (
    DocxCompiler,
    safe_autofit_table_to_contents,
    safe_set_cell_text,
    set_run_font_comprehensive,
)
from assessment_engine.domain.schemas.ast import (
    CellNode,
    DocumentAST,
    HeadingNode,
    TableNode,
    TableRowNode,
)
from assessment_engine.infrastructure.docx_render_utils import set_paragraph_numbering


def test_set_run_font_comprehensive_cjk_support():
    """Verify that set_run_font_comprehensive injects all 4 font quadrants (ASCII, hAnsi, CS, eastAsia)."""
    doc = docx.Document()
    p = doc.add_paragraph()
    run = p.add_run("テスト テキスト (Japanese Test Text)")

    set_run_font_comprehensive(run, font_name="Arial", east_asia_fallback="Meiryo")

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None
    assert rFonts.get(qn("w:ascii")) == "Arial"
    assert rFonts.get(qn("w:hAnsi")) == "Arial"
    assert rFonts.get(qn("w:cs")) == "Arial"
    assert rFonts.get(qn("w:eastAsia")) == "Meiryo"


def test_safe_set_cell_text_single_and_multiline():
    """Verify that safe_set_cell_text renders both single-line and multi-line content correctly."""
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)

    # Test single-line cell (e.g. heatmap header or metric)
    cell_single = table.rows[0].cells[0]
    safe_set_cell_text(cell_single, "TEF \\ LM")
    assert "TEF \\ LM" in cell_single.text

    # Test single-line with label: description pattern
    cell_label = table.rows[0].cells[1]
    safe_set_cell_text(cell_label, "Disponibilidad: 99.95%")
    assert "Disponibilidad:" in cell_label.text
    assert "99.95%" in cell_label.text

    # Test multi-line list cell
    cell_multi = table.rows[1].cells[0]
    safe_set_cell_text(cell_multi, "• Punto A\n• Punto B")
    assert "Punto A" in cell_multi.text
    assert "Punto B" in cell_multi.text


def test_safe_set_cell_text_empty_whitespace_newlines_guard():
    """Verify that empty lines or pure whitespace do not leave a cell with 0 paragraphs."""
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]

    safe_set_cell_text(cell, "\n\n   \n\t\n")
    # Must contain at least one paragraph to prevent MS Word fatal corruption
    assert len(cell.paragraphs) >= 1


def test_safe_autofit_table_to_contents_margin_boundary_and_xsd():
    """Verify that table width is bounded to margin percentage and respects CT_TblPr XSD sequence."""
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    tblPr = table._tbl.tblPr

    from assessment_engine.adapters.compilers.docx_compiler import _insert_in_tblpr

    # Insert borders in XSD schema order
    borders = OxmlElement("w:tblBorders")
    _insert_in_tblpr(tblPr, borders, "tblBorders")

    safe_autofit_table_to_contents(table)

    tblW = tblPr.find(qn("w:tblW"))
    assert tblW is not None
    assert tblW.get(qn("w:type")) == "pct"
    assert tblW.get(qn("w:w")) == "5000"

    # Verify order in CT_TblPr: tblW must precede tblBorders and tblBorders must precede tblLayout
    tags = [child.tag.split("}")[-1] for child in tblPr]
    assert "tblW" in tags
    assert "tblBorders" in tags
    assert "tblLayout" in tags
    assert tags.index("tblW") < tags.index("tblBorders")
    assert tags.index("tblBorders") < tags.index("tblLayout")


def test_set_paragraph_numbering_xsd_ordering():
    """Verify that w:numPr is inserted before w:spacing and w:ind complying with OpenXML CT_PPr schema."""
    doc = docx.Document()
    p = doc.add_paragraph()

    # Pre-add spacing element to simulate existing paragraph layout
    pPr = p._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "120")
    pPr.append(spacing)

    # Apply numbering
    success = set_paragraph_numbering(p, num_id=1, ilvl=0)
    assert success is True

    # Check child order: w:numPr must appear before w:spacing in CT_PPr
    tags = [child.tag.split("}")[-1] for child in pPr]
    assert "numPr" in tags
    assert "spacing" in tags
    num_idx = tags.index("numPr")
    spacing_idx = tags.index("spacing")
    assert num_idx < spacing_idx, f"numPr (index {num_idx}) must precede spacing (index {spacing_idx})"


def test_table_compilation_cant_split_and_tbl_header():
    """Verify that compiled table rows have w:cantSplit and header row has w:tblHeader."""
    compiler = DocxCompiler()
    doc = docx.Document()

    table_node = TableNode(
        rows=[
            TableRowNode(cells=[CellNode(text="Header 1"), CellNode(text="Header 2")]),
            TableRowNode(cells=[CellNode(text="Data 1"), CellNode(text="Data 2")]),
        ],
        autofit=True,
    )

    compiler._compile_table(doc, table_node)
    table = doc.tables[0]

    # Row 0 (Header): must have cantSplit AND tblHeader
    trPr_0 = table.rows[0]._tr.trPr
    assert trPr_0.find(qn("w:cantSplit")) is not None
    assert trPr_0.find(qn("w:tblHeader")) is not None

    # Row 1 (Data): must have cantSplit
    trPr_1 = table.rows[1]._tr.trPr
    assert trPr_1.find(qn("w:cantSplit")) is not None


def test_heading_compilation_keep_next_and_keep_lines():
    """Verify that compiled headings inject w:keepNext and w:keepLines to eliminate orphaned titles."""
    compiler = DocxCompiler()
    doc = docx.Document()

    heading_node = HeadingNode(text="1. Resumen Ejecutivo de Arquitectura", level=1)
    compiler._compile_heading(doc, heading_node)

    # Find compiled heading paragraph
    p = doc.paragraphs[-1]
    pPr = p._p.pPr
    assert pPr.find(qn("w:keepNext")) is not None
    assert pPr.find(qn("w:keepLines")) is not None
