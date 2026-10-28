"""Top-Tier-1 multi-format evidence ingestion and semantic parsing unit tests."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from assessment_engine.infrastructure.evidence_engine import EvidenceEngine


def test_markdown_parsing_and_heading_stack_hierarchy(tmp_path: Path):
    """Verify that _parse_text infers nested heading hierarchies from Markdown syntax."""
    storage_dir = tmp_path / "storage"
    engine = EvidenceEngine(client_id="test_client", storage_dir=storage_dir)

    md_content = """# Arquitectura Global
Esta es una descripción general de la infraestructura con más de veinte caracteres.

## Torre 1 Workplace
Puestos de trabajo virtuales desplegados en Azure con Windows 11 Enterprise.

### Conectividad WAN
Enlaces dedicados MPLS con respaldo 5G redundante para alta disponibilidad.
"""
    md_file = tmp_path / "architecture.md"
    md_file.write_text(md_content, encoding="utf-8")

    fragments = engine._parse_text(md_file)
    assert len(fragments) >= 5

    # Heading level 1
    h1_frag = [f for f in fragments if f.location_metadata.get("level") == 1][0]
    assert h1_frag.content == "Arquitectura Global"
    assert h1_frag.location_metadata["hierarchy"] == ["Arquitectura Global"]

    # Leaf line under level 3
    wan_frag = [f for f in fragments if "Enlaces dedicados MPLS" in f.content][0]
    assert wan_frag.location_metadata["hierarchy"] == [
        "Arquitectura Global",
        "Torre 1 Workplace",
        "Conectividad WAN",
    ]


def test_html_parsing_and_script_sanitization(tmp_path: Path):
    """Verify that _parse_html strips malicious scripts and extracts semantic text blocks."""
    storage_dir = tmp_path / "storage"
    engine = EvidenceEngine(client_id="test_client", storage_dir=storage_dir)

    html_content = """<!DOCTYPE html>
<html>
<head>
    <title>Informe de Auditoría Externa</title>
    <style>body { color: red; }</style>
    <script>alert("malicious_execution");</script>
</head>
<body>
    <h1>Resumen de Seguridad Perimetral</h1>
    <script>console.log("tracking");</script>
    <p>Se ha verificado la correcta implantación de cortafuegos de nueva generación NGFW en el perímetro exterior.</p>
    <li>Autenticación multifactor obligatoria en todos los accesos administrativos remotos VPN.</li>
</body>
</html>
"""
    html_file = tmp_path / "audit.html"
    html_file.write_text(html_content, encoding="utf-8")

    fragments = engine._parse_html(html_file)
    assert len(fragments) >= 3

    # Ensure scripts and styles were decomposed and never ingested as evidence
    all_content = " ".join(f.content for f in fragments)
    assert "malicious_execution" not in all_content
    assert "tracking" not in all_content
    assert "color: red" not in all_content

    # Ensure semantic paragraphs were extracted
    assert "cortafuegos de nueva generación NGFW" in all_content
    assert "Autenticación multifactor obligatoria" in all_content


def test_docx_parsing_paragraphs_and_tables(tmp_path: Path):
    """Verify that _parse_docx parses both running paragraphs and table cells into fragments."""
    storage_dir = tmp_path / "storage"
    engine = EvidenceEngine(client_id="test_client", storage_dir=storage_dir)

    doc = Document()
    doc.add_heading("Capítulo 1: Infraestructura de Cómputo", level=1)
    doc.add_paragraph("La plataforma actual cuenta con 120 hipervisores VMware vSphere en versión 7.0.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Componente"
    table.rows[0].cells[1].text = "Estado Operacional"
    table.rows[1].cells[0].text = "Cluster SAN Principal"
    table.rows[1].cells[1].text = "Saturación de almacenamiento al 88% de capacidad total."

    docx_file = tmp_path / "infra.docx"
    doc.save(str(docx_file))

    fragments = engine._parse_docx(docx_file)
    assert len(fragments) >= 2

    all_content = " ".join(f.content for f in fragments)
    assert "VMware vSphere" in all_content
    assert "Cluster SAN Principal" in all_content
    assert "Saturación de almacenamiento" in all_content


def test_get_grounding_context_and_ledger_persistence(tmp_path: Path):
    """Verify that get_grounding_context formats citations and ledger persists atomically."""
    storage_dir = tmp_path / "storage"
    engine = EvidenceEngine(client_id="test_client", storage_dir=storage_dir)

    sample_txt = tmp_path / "sample.txt"
    sample_txt.write_text(
        "Fragmento estratégico número uno con información relevante para la consultoría técnica.\n"
        "Fragmento estratégico número dos detallando SLAs y métricas cuantitativas del servicio.\n",
        encoding="utf-8",
    )

    ingested = engine.ingest_file(sample_txt)
    assert ingested >= 2

    # Verify ledger file on disk
    ledger_file = storage_dir / "evidence_vault.json"
    assert ledger_file.is_file()

    # Grounding context generation
    context = engine.get_grounding_context()
    assert len(context) > 50
    assert "[F-" in context or "Fragmento" in context
