"""Unit tests for W3C Design Tokens and Theme Registry (RFC-028 / Split-Brain #3)."""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml.ns import qn

from assessment_engine.adapters.emitters.docx_style_resolver import DocxStyleResolver
from assessment_engine.domain.theme_registry import (
    FontStackToken,
    HexColorToken,
    ScriptType,
    SemanticColorRole,
    ThemeRegistry,
    TypographyRole,
)


def test_hex_color_token_validation_and_rgb_tuple() -> None:
    """Verify HexColorToken format validation and RGB integer tuple extraction."""
    token = HexColorToken(value="#0072bc")
    assert token.value == "#0072BC"
    assert token.raw_hex == "0072BC"
    assert token.rgb_tuple == (0, 114, 188)

    # Invalid hex formats
    with pytest.raises(ValueError):
        HexColorToken(value="0072BC")  # Missing #

    with pytest.raises(ValueError):
        HexColorToken(value="#ZZZZZZ")  # Invalid hex characters

    with pytest.raises(ValueError):
        HexColorToken(value="#12345")  # Too short


def test_font_stack_resolution_by_script() -> None:
    """Verify internationalized font family matrix resolution across ISO 15924 scripts."""
    stack = FontStackToken(
        latin="Inter",
        cjk="Noto Sans CJK JP",
        arabic="Noto Sans Arabic",
        cyrillic="Inter",
        monospace="JetBrains Mono",
    )

    assert stack.resolve_for_script(ScriptType.LATIN) == "Inter"
    assert stack.resolve_for_script(ScriptType.CJK) == "Noto Sans CJK JP"
    assert stack.resolve_for_script(ScriptType.ARABIC) == "Noto Sans Arabic"
    assert stack.resolve_for_script(ScriptType.CYRILLIC) == "Inter"
    assert stack.resolve_for_script(ScriptType.LATIN, is_mono=True) == "JetBrains Mono"


def test_theme_registry_default_bootstrap_and_retrieval() -> None:
    """Verify default enterprise theme bootstrap and fallback retrieval."""
    theme = ThemeRegistry.get_theme()
    assert theme.theme_id == "global_enterprise_default"
    assert SemanticColorRole.BRAND_PRIMARY in theme.colors
    assert TypographyRole.HEADING_1 in theme.typography

    # Unknown theme returns default manifest
    fallback_theme = ThemeRegistry.get_theme("unknown_custom_tenant_theme")
    assert fallback_theme.theme_id == "global_enterprise_default"


def test_docx_style_resolver_typography_and_openxml_fonts() -> None:
    """Verify DocxStyleResolver applies semantic design tokens and w:rFonts matrix."""
    theme = ThemeRegistry.get_theme()
    resolver = DocxStyleResolver(theme, language_code="ja")  # Japanese CJK document

    assert resolver.script == ScriptType.CJK
    assert resolver.resolve_hex_string(SemanticColorRole.BRAND_PRIMARY) == "0072BC"
    assert resolver.resolve_rgb_tuple(SemanticColorRole.BRAND_PRIMARY) == (0, 114, 188)

    doc = Document()
    p = doc.add_paragraph()
    resolver.apply_paragraph_typography(p, TypographyRole.HEADING_1)
    assert p.paragraph_format.space_before.pt == 14.0
    assert p.paragraph_format.space_after.pt == 6.0

    run = p.add_run("Enterprise Resilience Assessment")
    resolver.apply_run_typography(run, TypographyRole.HEADING_1)

    assert run.font.bold is True
    assert run.font.size.pt == 16.0
    assert run.font.color.rgb == (0, 114, 188)

    # Verify actual OpenXML w:rFonts element attributes
    rPr = run._r.rPr
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None
    assert rFonts.get(qn("w:ascii")) == "Noto Sans CJK SC"
    assert rFonts.get(qn("w:eastAsia")) == "Noto Sans CJK SC"
